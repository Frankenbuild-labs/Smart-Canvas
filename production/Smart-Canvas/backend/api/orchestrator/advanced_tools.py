"""
Advanced Tools for the Metatron Orchestrator Agent
"""

import logging
import json
import re
from typing import Dict, Any, List, Optional
from datetime import datetime

from .dependencies import OrchestratorDependencies

# Configure logging
logger = logging.getLogger(__name__)

# Import tool dependencies
try:
    import wikipedia
    import requests
    import praw
    from youtube_transcript_api import YouTubeTranscriptApi
    from pytube import YouTube
    from PIL import Image
    import cv2
    from PyPDF2 import PdfReader
    from docx import Document
    import openpyxl
    TOOLS_AVAILABLE = True
except ImportError as e:
    logger.warning(f"⚠️ Some advanced tools not available: {e}")
    TOOLS_AVAILABLE = False


async def wikipedia_search_tool(
    deps: OrchestratorDependencies,
    query: str,
    language: str = "en",
    sentences: int = 3
) -> str:
    """
    Search Wikipedia for comprehensive information
    """
    try:
        logger.info(f"📚 Wikipedia Search: {query}")
        
        if not TOOLS_AVAILABLE:
            return "Wikipedia search not available. Please install wikipedia package."
        
        wikipedia.set_lang(language)
        
        # Search for pages
        search_results = wikipedia.search(query, results=5)
        if not search_results:
            return f"No Wikipedia results found for '{query}'"
        
        # Get summary of the first result
        try:
            page = wikipedia.page(search_results[0])
            summary = wikipedia.summary(search_results[0], sentences=sentences)
            
            result = f"✅ Wikipedia: {page.title}\n\n"
            result += f"📄 Summary:\n{summary}\n\n"
            result += f"🔗 URL: {page.url}\n\n"
            
            if len(search_results) > 1:
                result += f"📋 Related articles: {', '.join(search_results[1:])}"
            
            return result
            
        except wikipedia.exceptions.DisambiguationError as e:
            # Handle disambiguation
            options = e.options[:5]
            return f"Multiple Wikipedia pages found for '{query}':\n" + "\n".join(f"• {option}" for option in options)
        
        except wikipedia.exceptions.PageError:
            return f"Wikipedia page not found for '{query}'"
        
    except Exception as e:
        logger.error(f"❌ Wikipedia Search error: {str(e)}")
        return f"Wikipedia search failed: {str(e)}"


async def reddit_search_tool(
    deps: OrchestratorDependencies,
    query: str,
    subreddit: str = None,
    sort: str = "relevance",
    limit: int = 10
) -> str:
    """
    Search Reddit for discussions and community insights
    """
    try:
        logger.info(f"🔴 Reddit Search: {query}")
        
        if not TOOLS_AVAILABLE:
            return "Reddit search not available. Please install praw package."
        
        # Use requests to search Reddit via JSON API (no credentials needed)
        try:
            import requests

            # Use Reddit's JSON API for public search
            if subreddit:
                search_url = f"https://www.reddit.com/r/{subreddit}/search.json"
                params = {
                    'q': query,
                    'restrict_sr': 'on',
                    'sort': sort,
                    'limit': min(limit, 25)
                }
            else:
                search_url = "https://www.reddit.com/search.json"
                params = {
                    'q': query,
                    'sort': sort,
                    'limit': min(limit, 25)
                }

            headers = {
                'User-Agent': 'Metatron-AI-Research/1.0'
            }

            response = requests.get(search_url, params=params, headers=headers, timeout=10)

            if response.status_code == 200:
                data = response.json()
                posts = data.get('data', {}).get('children', [])

                if not posts:
                    return f"❌ No Reddit posts found for '{query}'"

                result = f"✅ **Reddit Search Results for '{query}':**\n\n"

                if subreddit:
                    result += f"📍 Searching in r/{subreddit}\n\n"

                for i, post_data in enumerate(posts[:limit], 1):
                    post = post_data.get('data', {})

                    title = post.get('title', 'No title')
                    score = post.get('score', 0)
                    num_comments = post.get('num_comments', 0)
                    subreddit_name = post.get('subreddit', 'unknown')
                    permalink = post.get('permalink', '')
                    url = f"https://reddit.com{permalink}" if permalink else "No URL"

                    # Get post preview/selftext
                    selftext = post.get('selftext', '')
                    preview = selftext[:200] + "..." if len(selftext) > 200 else selftext

                    result += f"**{i}. {title}**\n"
                    result += f"📊 {score} upvotes, {num_comments} comments\n"
                    result += f"📍 r/{subreddit_name}\n"
                    if preview:
                        result += f"📝 {preview}\n"
                    result += f"🔗 {url}\n\n"

                return result

            else:
                return f"❌ Reddit API error: {response.status_code}"

        except requests.RequestException as e:
            return f"❌ Reddit search failed: Network error - {str(e)}"
        except Exception as e:
            return f"❌ Reddit search failed: {str(e)}"
        
        return result
        
    except Exception as e:
        logger.error(f"❌ Reddit Search error: {str(e)}")
        return f"Reddit search failed: {str(e)}"


async def news_research_tool(
    deps: OrchestratorDependencies,
    query: str,
    category: str = None,
    language: str = "en",
    limit: int = 10
) -> str:
    """
    Get latest news and information on specific topics
    """
    try:
        logger.info(f"📰 News Research: {query}")
        
        # Use NewsAPI.org free tier or web scraping for news
        try:
            import requests
            from datetime import datetime, timedelta

            # Try NewsAPI first (requires API key in production)
            # For now, use a free alternative - RSS feeds or web scraping

            # Use Google News RSS feed as a free alternative
            try:
                # Google News RSS feed
                rss_url = f"https://news.google.com/rss/search"
                params = {
                    'q': query,
                    'hl': language,
                    'gl': 'US',
                    'ceid': 'US:en'
                }

                headers = {
                    'User-Agent': 'Metatron-AI-Research/1.0'
                }

                response = requests.get(rss_url, params=params, headers=headers, timeout=10)

                if response.status_code == 200:
                    # Parse RSS feed
                    from xml.etree import ElementTree as ET

                    root = ET.fromstring(response.content)
                    items = root.findall('.//item')

                    if not items:
                        return f"❌ No news articles found for '{query}'"

                    result = f"✅ **Latest News for '{query}':**\n\n"

                    if category:
                        result += f"📂 Category: {category}\n\n"

                    for i, item in enumerate(items[:limit], 1):
                        title = item.find('title').text if item.find('title') is not None else "No title"
                        link = item.find('link').text if item.find('link') is not None else "No link"
                        pub_date = item.find('pubDate').text if item.find('pubDate') is not None else "No date"
                        description = item.find('description').text if item.find('description') is not None else ""

                        # Clean up description (remove HTML tags)
                        if description:
                            import re
                            description = re.sub('<[^<]+?>', '', description)
                            description = description[:200] + "..." if len(description) > 200 else description

                        result += f"**{i}. {title}**\n"
                        if description:
                            result += f"📝 {description}\n"
                        result += f"📅 {pub_date}\n"
                        result += f"🔗 {link}\n\n"

                    return result

                else:
                    # Fallback to a simple news search
                    return f"❌ News API error: {response.status_code}. Try using a more specific search term."

            except Exception as rss_error:
                logger.warning(f"RSS feed failed: {rss_error}")

                # Fallback: Return a helpful message about news search
                result = f"📰 **News Search for '{query}':**\n\n"
                result += f"🔍 To get real-time news results, please:\n"
                result += f"1. Configure a news API key (NewsAPI.org)\n"
                result += f"2. Or try searching directly on news websites\n"
                result += f"3. Suggested sources for '{query}':\n"
                result += f"   • Google News: https://news.google.com/search?q={query.replace(' ', '%20')}\n"
                result += f"   • Reuters: https://www.reuters.com/search/news?blob={query.replace(' ', '%20')}\n"
                result += f"   • BBC News: https://www.bbc.com/search?q={query.replace(' ', '%20')}\n\n"
                result += f"💡 Real-time news integration available with API configuration"

                return result

        except Exception as e:
            return f"❌ News search failed: {str(e)}"
        
        return result
        
    except Exception as e:
        logger.error(f"❌ News Research error: {str(e)}")
        return f"News research failed: {str(e)}"


async def document_processing_tool(
    deps: OrchestratorDependencies,
    file_path: str,
    action: str = "extract_text",
    parameters: Dict[str, Any] = None
) -> str:
    """
    Process documents (PDF, Word, Excel) and extract information
    """
    try:
        logger.info(f"📄 Document Processing: {file_path}")
        
        if not TOOLS_AVAILABLE:
            return "Document processing tools not available. Please install required packages."
        
        file_extension = file_path.lower().split('.')[-1]
        
        if action == "extract_text":
            if file_extension == "pdf":
                # PDF processing
                try:
                    reader = PdfReader(file_path)
                    text = ""
                    for page in reader.pages:
                        text += page.extract_text()
                    return f"✅ PDF text extracted from {file_path}:\n\n{text[:1500]}..."
                except Exception as e:
                    return f"PDF processing failed: {str(e)}"
            
            elif file_extension in ["docx", "doc"]:
                # Word document processing
                try:
                    doc = Document(file_path)
                    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                    return f"✅ Word document text extracted from {file_path}:\n\n{text[:1500]}..."
                except Exception as e:
                    return f"Word document processing failed: {str(e)}"
            
            elif file_extension in ["xlsx", "xls"]:
                # Excel processing
                try:
                    workbook = openpyxl.load_workbook(file_path)
                    sheet = workbook.active
                    data = []
                    for row in sheet.iter_rows(values_only=True):
                        data.append(row)
                    return f"✅ Excel data extracted from {file_path}:\n\n{str(data[:10])}..."
                except Exception as e:
                    return f"Excel processing failed: {str(e)}"
            
            else:
                return f"Unsupported file type: {file_extension}"
        
        else:
            return f"Unsupported action: {action}"
        
    except Exception as e:
        logger.error(f"❌ Document Processing error: {str(e)}")
        return f"Document processing failed: {str(e)}"


async def youtube_analysis_tool(
    deps: OrchestratorDependencies,
    url: str,
    include_transcript: bool = True,
    include_metadata: bool = True
) -> str:
    """
    Analyze YouTube videos, extract transcripts and metadata
    """
    try:
        logger.info(f"🎥 YouTube Analysis: {url}")
        
        if not TOOLS_AVAILABLE:
            return "YouTube analysis tools not available. Please install required packages."
        
        # Extract video ID from URL
        video_id_match = re.search(r'(?:v=|\/)([0-9A-Za-z_-]{11}).*', url)
        if not video_id_match:
            return "Invalid YouTube URL format"
        
        video_id = video_id_match.group(1)
        result = f"✅ YouTube Video Analysis: {url}\n\n"
        
        if include_metadata:
            try:
                yt = YouTube(url)
                result += f"📹 **Title**: {yt.title}\n"
                result += f"👤 **Channel**: {yt.author}\n"
                result += f"📅 **Published**: {yt.publish_date}\n"
                result += f"👀 **Views**: {yt.views:,}\n"
                result += f"⏱️ **Duration**: {yt.length} seconds\n"
                result += f"📝 **Description**: {yt.description[:200]}...\n\n"
            except Exception as e:
                result += f"⚠️ Metadata extraction failed: {str(e)}\n\n"
        
        if include_transcript:
            try:
                transcript = YouTubeTranscriptApi.get_transcript(video_id)
                transcript_text = " ".join([entry['text'] for entry in transcript])
                result += f"📜 **Transcript**:\n{transcript_text[:1000]}...\n\n"
            except Exception as e:
                result += f"⚠️ Transcript extraction failed: {str(e)}\n\n"
        
        return result

    except Exception as e:
        logger.error(f"❌ YouTube Analysis error: {str(e)}")
        return f"YouTube analysis failed: {str(e)}"


async def image_video_understanding_tool(
    deps: OrchestratorDependencies,
    media_path: str,
    media_type: str,
    analysis_type: str = "comprehensive"
) -> str:
    """
    Understand and analyze images, videos, and other media content
    """
    try:
        logger.info(f"🖼️ Media Understanding: {media_path} ({media_type})")

        if not TOOLS_AVAILABLE:
            return "Media understanding tools not available. Please install required packages."

        result = f"✅ Media Analysis: {media_path}\n\n"

        if media_type.lower() in ["jpg", "jpeg", "png", "bmp", "gif"]:
            # Image analysis
            try:
                img = Image.open(media_path)
                result += f"📐 **Dimensions**: {img.size[0]} x {img.size[1]} pixels\n"
                result += f"🎨 **Mode**: {img.mode}\n"
                result += f"📊 **Format**: {img.format}\n\n"

                if analysis_type == "comprehensive":
                    # Basic image analysis
                    result += f"🔍 **Analysis**: Image successfully loaded and analyzed\n"
                    result += f"💡 **Note**: Advanced AI vision analysis requires additional setup\n\n"

            except Exception as e:
                result += f"⚠️ Image analysis failed: {str(e)}\n\n"

        elif media_type.lower() in ["mp4", "avi", "mov", "mkv"]:
            # Video analysis
            try:
                cap = cv2.VideoCapture(media_path)
                frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
                fps = cap.get(cv2.CAP_PROP_FPS)
                duration = frame_count / fps if fps > 0 else 0
                width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
                height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))

                result += f"📹 **Duration**: {duration:.2f} seconds\n"
                result += f"📐 **Resolution**: {width} x {height}\n"
                result += f"🎬 **Frame Rate**: {fps:.2f} FPS\n"
                result += f"📊 **Total Frames**: {frame_count}\n\n"

                cap.release()

            except Exception as e:
                result += f"⚠️ Video analysis failed: {str(e)}\n\n"

        else:
            result += f"⚠️ Unsupported media type: {media_type}\n\n"

        return result

    except Exception as e:
        logger.error(f"❌ Media Understanding error: {str(e)}")
        return f"Media understanding failed: {str(e)}"


async def mermaid_diagram_tool(
    deps: OrchestratorDependencies,
    diagram_type: str,
    content: str,
    title: str = None
) -> str:
    """
    Generate Mermaid diagrams for flowcharts, sequences, and visualizations
    """
    try:
        logger.info(f"📊 Mermaid Diagram: {diagram_type}")

        # Mermaid diagram templates
        templates = {
            "flowchart": """```mermaid
flowchart TD
    {content}
```""",
            "sequence": """```mermaid
sequenceDiagram
    {content}
```""",
            "gantt": """```mermaid
gantt
    title {title}
    {content}
```""",
            "pie": """```mermaid
pie title {title}
    {content}
```""",
            "gitgraph": """```mermaid
gitGraph
    {content}
```""",
            "mindmap": """```mermaid
mindmap
  root{{{title}}}
    {content}
```""",
            "timeline": """```mermaid
timeline
    title {title}
    {content}
```"""
        }

        if diagram_type not in templates:
            available_types = ", ".join(templates.keys())
            return f"❌ Unsupported diagram type: {diagram_type}\nAvailable types: {available_types}"

        # Generate the diagram
        template = templates[diagram_type]
        diagram = template.format(
            content=content,
            title=title or f"{diagram_type.title()} Diagram"
        )

        result = f"✅ Mermaid {diagram_type.title()} Diagram Generated:\n\n"
        result += diagram
        result += f"\n\n💡 **Usage**: Copy the above code block and paste it in any Markdown-compatible editor or the frontend chat interface."

        return result

    except Exception as e:
        logger.error(f"❌ Mermaid Diagram error: {str(e)}")
        return f"Mermaid diagram generation failed: {str(e)}"
